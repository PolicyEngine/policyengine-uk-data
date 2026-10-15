"""Hash-bound OBR workbooks, header-read years and required target coverage.

The OBR targets come only from the workbook bytes the active manifest names.
These tests pin that:
- other bytes, or a workbook with another title, are refused;
- year columns follow the header row, not the column letter;
- a missing target or year fails;
- ingesting a vintage reproduces the build's own targets.
"""

import io
import shutil

import openpyxl
import pytest
from hypothesis import HealthCheck, given, settings
from hypothesis import strategies as st

from policyengine_uk_data.storage import STORAGE_FOLDER
from policyengine_uk_data.targets import obr_sources
from policyengine_uk_data.targets.obr_sources import (
    SourceError,
    active_manifest,
    check_coverage,
    coverage_problems,
    ingest,
    load_workbook,
    parse_manifest,
    read_row,
    verified_bytes,
    year_columns,
)
from policyengine_uk_data.targets.sources import obr

MANIFEST = active_manifest()


def _workbook_targets(receipts, expenditure, manifest=MANIFEST):
    return [
        *obr._parse_receipts(receipts, manifest),
        *obr._parse_nics(receipts, manifest),
        *obr._parse_council_tax(expenditure, manifest),
        *obr._parse_welfare(expenditure, manifest),
        *obr._parse_tv_licence(expenditure, manifest),
    ]


def _copy_storage(tmp_path):
    for entry in MANIFEST.files.values():
        target = tmp_path / entry.path
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copyfile(STORAGE_FOLDER / entry.path, target)
    return tmp_path


def test_the_build_reads_the_manifest_files_and_names_their_vintage():
    assert MANIFEST.vintage == "march_2026"
    for kind, entry in MANIFEST.files.items():
        assert verified_bytes(entry)
        assert obr_sources.workbook_title(load_workbook(kind)) == entry.title
    targets = obr.get_targets()
    built = [t for t in targets if t.forecast_vintage == MANIFEST.vintage]
    assert len(built) == 25
    assert {t.reference_url for t in built} == {f.url for f in MANIFEST.files.values()}
    check_coverage(built, MANIFEST)


@settings(
    max_examples=25,
    deadline=None,
    suppress_health_check=[HealthCheck.function_scoped_fixture],
)
@given(data=st.data())
def test_any_change_to_the_bytes_is_refused(tmp_path, data):
    storage = _copy_storage(tmp_path)
    entry = MANIFEST.files[data.draw(st.sampled_from(["receipts", "expenditure"]))]
    path = storage / entry.path
    original = (STORAGE_FOLDER / entry.path).read_bytes()
    kind = data.draw(st.sampled_from(["flip", "truncate", "append"]))
    if kind == "flip":
        i = data.draw(st.integers(0, len(original) - 1))
        changed = original[:i] + bytes([original[i] ^ 0x01]) + original[i + 1 :]
    elif kind == "truncate":
        changed = original[: data.draw(st.integers(0, len(original) - 1))]
    else:
        changed = original + data.draw(st.binary(min_size=1, max_size=16))
    path.write_bytes(changed)
    with pytest.raises(SourceError, match="not the file its manifest names"):
        verified_bytes(entry, storage)


def test_a_missing_workbook_fails_the_targets(tmp_path):
    with pytest.raises(SourceError, match="missing"):
        obr.workbook_targets(MANIFEST, tmp_path)


def test_a_workbook_titled_for_another_vintage_is_refused(tmp_path):
    """Bytes that match a manifest but carry another vintage's title fail."""
    entry = MANIFEST.files["receipts"]
    wb = openpyxl.load_workbook(STORAGE_FOLDER / entry.path)
    wb["Contents"]["B2"] = entry.title.replace("March 2026", "November 2025")
    buffer = io.BytesIO()
    wb.save(buffer)
    data = buffer.getvalue()
    (tmp_path / "x.xlsx").write_bytes(data)
    relabelled = obr_sources.SourceFile(
        kind="receipts",
        path="x.xlsx",
        sha256=obr_sources.sha256_bytes(data),
        size_bytes=len(data),
        title=entry.title,
        url=entry.url,
        retrieved="test",
    )
    manifest = obr_sources.ObrManifest(
        **{**MANIFEST.__dict__, "files": {**MANIFEST.files, "receipts": relabelled}}
    )
    with pytest.raises(SourceError, match="November 2025"):
        load_workbook("receipts", manifest, tmp_path)


# The March 2026 tables' year columns as the loader hard-coded them before it
# read the headers: an independent reader for a differential check.
_FIXED = {
    "3.4": dict(zip("CDEFGHI", range(2024, 2031))),
    "3.8": dict(zip("DEFGHIJ", range(2024, 2031))),
    "4.1": dict(zip("CDEFGHI", range(2024, 2031))),
    "4.9": dict(zip("CDEFGHI", range(2024, 2031))),
    "4.19": dict(zip("CDEFGHI", range(2024, 2031))),
}
_ROWS = {
    "obr/income_tax": ("receipts", "3.4", "Income tax (gross of tax credits)"),
    "obr/ni_employer": ("receipts", "3.4", "Class 1 Employer NICs"),
    "obr/vat": ("receipts", "3.8", "Value added tax"),
    "obr/sdlt": ("receipts", "3.8", "Stamp duty land tax"),
    "obr/council_tax_wales": ("expenditure", "4.1", "Wales council tax receipts"),
    "obr/child_benefit": ("expenditure", "4.9", "Child benefit"),
    "obr/state_pension": ("expenditure", "4.9", "State pension"),
}


def test_header_read_years_match_the_fixed_columns_on_march_2026():
    books = {kind: load_workbook(kind) for kind in ("receipts", "expenditure")}
    targets = {
        t.name: t for t in _workbook_targets(books["receipts"], books["expenditure"])
    }
    for sheet, fixed in _FIXED.items():
        ws = books["receipts" if sheet.startswith("3") else "expenditure"][sheet]
        assert year_columns(ws, 2024, 2030) == fixed, sheet
    for name, (kind, sheet, label) in _ROWS.items():
        ws = books[kind][sheet]
        row = obr._find_row(ws, label, max_row=80)
        expected = {
            year: ws[f"{c}{row}"].value * 1e9 for c, year in _FIXED[sheet].items()
        }
        assert targets[name].values == pytest.approx(expected), name


def _sheet(header_row, first_column, years, values_by_year, unit="£ billion"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = unit
    for offset, year in enumerate(years):
        column = first_column + offset
        ws.cell(row=header_row, column=column, value=f"{year}-{(year + 1) % 100:02d}")
        ws.cell(row=12, column=column, value=values_by_year[year])
    return ws


@settings(max_examples=200, deadline=None)
@given(
    header_row=st.integers(2, 10),
    first_column=st.integers(2, 6),
    first_year=st.integers(2020, 2030),
    length=st.integers(3, 8),
    data=st.data(),
)
def test_values_follow_their_header_whatever_the_column_order(
    header_row, first_column, first_year, length, data
):
    years = list(range(first_year, first_year + length))
    order = data.draw(st.permutations(years))
    values = {year: float(year - 2000) for year in years}
    ws = _sheet(header_row, first_column, order, values)
    manifest = obr_sources.ObrManifest(
        **{**MANIFEST.__dict__, "first_year": years[0], "last_year": years[-1]}
    )
    assert read_row(ws, 12, manifest) == {y: v * 1e9 for y, v in values.items()}
    # Declaring the vintage a year later or earlier than the table fails.
    for shift in (-1, 1):
        with pytest.raises(SourceError, match="the manifest declares"):
            year_columns(ws, years[0] + shift, years[-1] + shift)


def test_a_table_without_headers_or_unit_fails():
    ws = openpyxl.Workbook().active
    ws["C12"] = 1.0
    with pytest.raises(SourceError, match="financial-year headers"):
        year_columns(ws, 2024, 2030)
    ws = _sheet(
        5, 3, range(2024, 2031), {y: 1.0 for y in range(2024, 2031)}, unit="£ million"
    )
    with pytest.raises(SourceError, match="unit label"):
        read_row(ws, 12, MANIFEST)


def test_a_repeated_or_malformed_year_header_fails():
    ws = _sheet(5, 3, [2024, 2025, 2025], {2024: 1.0, 2025: 1.0})
    with pytest.raises(SourceError, match="heads two columns"):
        year_columns(ws, 2024, 2025)
    ws = openpyxl.Workbook().active
    ws["C5"] = "2024-26"
    with pytest.raises(SourceError, match="not one financial year"):
        year_columns(ws, 2024, 2024)


@pytest.fixture(scope="module")
def march_targets():
    return obr.workbook_targets()


@settings(
    max_examples=60,
    deadline=None,
    suppress_health_check=[HealthCheck.function_scoped_fixture],
)
@given(data=st.data())
def test_losing_any_year_of_any_target_fails_coverage(march_targets, data):
    targets = [t for t in march_targets if t.forecast_vintage == MANIFEST.vintage]
    victim = data.draw(st.sampled_from(targets))
    year = data.draw(st.sampled_from(sorted(victim.values)))
    damaged = [
        t.model_copy(
            update={"values": {y: v for y, v in t.values.items() if y != year}}
        )
        if t is victim
        else t
        for t in targets
    ]
    problems = coverage_problems(damaged, MANIFEST)
    assert problems == [f"{victim.name} has no value for {year}"]


@settings(
    max_examples=40,
    deadline=None,
    suppress_health_check=[HealthCheck.function_scoped_fixture],
)
@given(data=st.data())
def test_losing_any_required_target_fails_coverage(march_targets, data):
    name = data.draw(st.sampled_from(MANIFEST.required))
    problems = coverage_problems([t for t in march_targets if t.name != name], MANIFEST)
    assert problems == [f"{name} was not built"]


def test_a_layout_change_in_the_workbook_fails_the_build():
    """A renamed row, a formula in a value cell, or a lost self-employed line."""
    receipts, expenditure = load_workbook("receipts"), load_workbook("expenditure")
    ws = receipts[obr._find_receipts_sheet(receipts).title]
    ws[f"B{obr._find_row(ws, 'Value added tax', max_row=80)}"] = "VAT (renamed)"
    ws49 = expenditure["4.9"]
    ws49[f"F{obr._find_row(ws49, 'Child benefit', max_row=55)}"] = "=1+1"
    ws34 = receipts["3.4"]
    ws34[
        f"B{obr._find_row(ws34, 'Class 4 and Class 2 Self employed NICs', max_row=40)}"
    ] = "x"
    problems = coverage_problems(_workbook_targets(receipts, expenditure), MANIFEST)
    assert "obr/vat was not built" in problems
    assert "obr/child_benefit has no value for 2027" in problems
    assert (
        "none of obr/ni_self_employed or obr/ni_class_2+obr/ni_class_4 was built"
        in problems
    )


def test_a_target_missing_from_the_manifest_fails(march_targets):
    manifest = obr_sources.ObrManifest(
        **{
            **MANIFEST.__dict__,
            "required": tuple(
                n for n in MANIFEST.required if n != "obr/tv_licence_fee"
            ),
        }
    )
    assert coverage_problems(march_targets, manifest) == [
        "obr/tv_licence_fee was built but is not in the manifest"
    ]


def _sources(tmp_path):
    src = tmp_path / "downloads"
    src.mkdir(exist_ok=True)
    out = {}
    for kind, entry in MANIFEST.files.items():
        out[kind] = src / f"{kind}.xlsx"
        shutil.copyfile(STORAGE_FOLDER / entry.path, out[kind])
    return out


def _ingest(tmp_path, **overrides):
    (tmp_path / "manifests").mkdir(exist_ok=True)
    arguments = dict(
        vintage="march_2026",
        label="March 2026",
        sources=overrides.pop("sources", None) or _sources(tmp_path),
        urls={k: f.url for k, f in MANIFEST.files.items()},
        first_year=2024,
        last_year=2030,
        current=MANIFEST,
        storage=tmp_path / "storage",
        manifest_dir=tmp_path / "manifests",
        retrieved="test",
    )
    arguments.update(overrides)
    return ingest(**arguments)


def test_ingesting_the_same_workbooks_reproduces_the_targets(tmp_path, march_targets):
    manifest, path, diff = _ingest(tmp_path)
    assert path.name == "obr_march_2026.yaml"
    assert manifest.vintage == "march_2026"
    assert {k: f.sha256 for k, f in manifest.files.items()} == {
        k: f.sha256 for k, f in MANIFEST.files.items()
    }
    assert manifest.files["receipts"].path == "obr_efo/march_2026/receipts.xlsx"
    assert diff and all(row["old"] == row["new"] for row in diff)
    again = obr.workbook_targets(manifest, tmp_path / "storage")
    assert {t.name: t.values for t in again} == {
        t.name: t.values for t in march_targets
    }
    assert {t.forecast_vintage for t in again} == {"march_2026"}


def test_ingest_refuses_a_workbook_that_does_not_name_the_vintage(tmp_path):
    """Sol review of #592, P1: March bytes cannot be adopted as October."""
    with pytest.raises(SourceError, match="not the October 2026 EFO receipts"):
        _ingest(tmp_path, vintage="october_2026", label="October 2026")
    with pytest.raises(SourceError, match="does not name vintage"):
        _ingest(tmp_path, vintage="october_2026", label="March 2026")


def test_ingest_refuses_years_the_tables_do_not_head_and_keeps_nothing(tmp_path):
    with pytest.raises(SourceError, match="the manifest declares 2025 to 2031"):
        _ingest(tmp_path, first_year=2025, last_year=2031)
    assert not (tmp_path / "storage" / "obr_efo" / "march_2026").exists()
    assert not list((tmp_path / "manifests").iterdir())


def test_ingest_refuses_the_page_obr_serves_to_scripts(tmp_path):
    sources = _sources(tmp_path)
    sources["receipts"].write_bytes(
        b"<!DOCTYPE html><html><head><title>No Access</title></head></html>"
    )
    with pytest.raises(SourceError, match="download the workbook in a browser"):
        _ingest(tmp_path, sources=sources)


def test_verify_command_passes_on_the_active_manifest(capsys):
    assert obr_sources.main(["verify"]) == 0
    assert "both workbooks verified, 25 targets" in capsys.readouterr().out


@pytest.mark.parametrize(
    "change, match",
    [
        ({"schema_version": 2}, "schema_version"),
        ({"unit": "£ million"}, "unit"),
        ({"first_year": 2031}, "first_year"),
        ({"targets": {"required": ["obr/vat", "obr/vat"]}}, "twice"),
        ({"targets": {"required": []}}, "no required"),
    ],
)
def test_malformed_manifests_are_refused(change, match):
    import yaml

    raw = yaml.safe_load(MANIFEST.path.read_text())
    raw.update(change)
    with pytest.raises(SourceError, match=match):
        parse_manifest(raw)


def test_a_bad_hash_in_the_manifest_is_refused():
    import yaml

    raw = yaml.safe_load(MANIFEST.path.read_text())
    raw["files"]["receipts"]["sha256"] = "abc"
    with pytest.raises(SourceError, match="64 hex"):
        parse_manifest(raw)


def test_a_manifest_cannot_relabel_verified_bytes():
    """Sol review of #592, P1: the vintage, label and titles must agree."""
    import yaml

    raw = yaml.safe_load(MANIFEST.path.read_text())
    raw.update(vintage="october_2026", label="October 2026")
    with pytest.raises(SourceError, match="not the October 2026 EFO"):
        parse_manifest(raw)
    raw = yaml.safe_load(MANIFEST.path.read_text())
    raw.update(label="October 2026")
    with pytest.raises(SourceError, match="does not name vintage"):
        parse_manifest(raw)


def test_a_row_stating_its_own_unit_is_refused():
    """Sol review of #592, P1: "£ billion unless otherwise stated" tables."""
    expenditure = load_workbook("expenditure")
    ws = expenditure["4.19"]
    row = obr._find_row(ws, "Standard colour licence fee (£)", max_row=30)
    with pytest.raises(SourceError, match="states its own unit"):
        read_row(ws, row, MANIFEST)
    # A renamed receipts row no longer falls through to the fee row.
    ws[f"B{obr._find_row(ws, 'BBC licence fee receipts', max_row=30)}"] = (
        "BBC television licensing receipts"
    )
    receipts = load_workbook("receipts")
    problems = coverage_problems(_workbook_targets(receipts, expenditure), MANIFEST)
    assert problems == ["obr/tv_licence_fee was not built"]


def test_ingest_never_replaces_an_ingested_vintage(tmp_path):
    """Sol review of #592, P2: a retry cannot delete an adopted vintage."""
    _ingest(tmp_path)
    kept = (
        tmp_path / "storage" / "obr_efo" / "march_2026" / "receipts.xlsx"
    ).read_bytes()
    with pytest.raises(SourceError, match="already ingested"):
        _ingest(tmp_path, first_year=2025, last_year=2031)
    assert (
        tmp_path / "storage" / "obr_efo" / "march_2026" / "receipts.xlsx"
    ).read_bytes() == kept
    assert (tmp_path / "manifests" / "obr_march_2026.yaml").exists()


def test_ingest_checks_both_workbooks_before_copying_either(tmp_path):
    sources = _sources(tmp_path)
    sources["expenditure"].write_bytes(b"not a workbook")
    with pytest.raises(SourceError, match="not an xlsx workbook"):
        _ingest(tmp_path, sources=sources)
    assert not (tmp_path / "storage").exists()


def test_a_workbook_target_without_the_vintage_fails_coverage(march_targets):
    """Sol review of #592, P2: an unversioned target cannot skip the checks."""
    extra = march_targets[0].model_copy(
        update={
            "name": "obr/unversioned",
            "forecast_vintage": None,
            "values": {2024: 1.0},
        }
    )
    problems = coverage_problems(list(march_targets) + [extra], MANIFEST)
    assert "obr/unversioned carries vintage None, not 'march_2026'" in problems
    assert "obr/unversioned was built but is not in the manifest" in problems


def test_ingest_records_each_download_time(tmp_path):
    """A browser download is hash-bound and timed: the file's modification
    time by default, or the time given."""
    import os

    sources = _sources(tmp_path)
    from datetime import datetime, timezone

    when = datetime(2026, 10, 18, 12, tzinfo=timezone.utc).timestamp()
    for path in sources.values():
        os.utime(path, (when, when))
    manifest, path, _ = _ingest(tmp_path, sources=sources)
    assert {f.downloaded_at for f in manifest.files.values()} == {
        "2026-10-18T12:00:00+00:00"
    }
    assert "downloaded_at: '2026-10-18T12:00:00+00:00'" in path.read_text()
    (tmp_path / "again").mkdir()
    manifest, _, _ = _ingest(
        tmp_path / "again",
        downloaded_at={
            "receipts": "2026-10-28T13:05:00+00:00",
            "expenditure": "2026-10-28T13:06:00+00:00",
        },
    )
    assert manifest.files["receipts"].downloaded_at == "2026-10-28T13:05:00+00:00"
    with pytest.raises(SourceError, match="UTC offset"):
        _ingest(
            tmp_path / "again",
            vintage="march_2026",
            downloaded_at={
                "receipts": "2026-10-28T13:05:00",
                "expenditure": "2026-10-28T13:06:00+00:00",
            },
            dry_run=True,
        )


def test_a_dry_run_checks_everything_and_writes_nothing(tmp_path):
    manifest, path, diff = _ingest(tmp_path, dry_run=True)
    assert manifest.vintage == "march_2026"
    assert diff and all(row["old"] == row["new"] for row in diff)
    assert not path.exists()
    assert not (tmp_path / "storage" / "obr_efo").exists()
    # A dry run still refuses workbooks that do not carry the vintage.
    with pytest.raises(SourceError, match="not the October 2026 EFO receipts"):
        _ingest(tmp_path, vintage="october_2026", label="October 2026", dry_run=True)


def test_the_rehearsal_command_runs_on_the_committed_march_workbooks(tmp_path, capsys):
    """The Budget-day command, end to end, on downloads of the active vintage."""
    sources = _sources(tmp_path)
    code = obr_sources.main(
        [
            "ingest",
            "--vintage",
            "march_2026",
            "--label",
            "March 2026",
            "--receipts",
            str(sources["receipts"]),
            "--expenditure",
            str(sources["expenditure"]),
            "--first-year",
            "2024",
            "--last-year",
            "2030",
            "--dry-run",
        ]
    )
    out = capsys.readouterr().out
    assert code == 0
    assert "the same bytes as the committed March 2026 workbook" in out
    assert "would write" in out
