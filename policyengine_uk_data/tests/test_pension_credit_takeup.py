"""Pension Credit take-up solved over entitled units, and DWP targets."""

import numpy as np
import openpyxl
from hypothesis import given, settings
from hypothesis import strategies as st

from policyengine_uk_data.parameters import load_take_up_rate
from policyengine_uk_data.targets import get_all_targets
from policyengine_uk_data.targets.schema import GREAT_BRITAIN, Unit
from policyengine_uk_data.targets.sources import dwp_pension_credit, obr
from policyengine_uk_data.datasets.pension_credit_takeup import (
    pension_credit_takeup_flags,
)
from policyengine_uk_data.utils.takeup import solve_fill_probability

_units = st.lists(
    st.tuples(st.floats(0, 5_000), st.booleans(), st.booleans()),
    min_size=1,
    max_size=60,
)


@settings(max_examples=300, deadline=None)
@given(_units, st.floats(0, 1))
def test_fill_probability_meets_the_rate_or_hits_a_bound(units, rate):
    weights = np.array([u[0] for u in units])
    eligible = np.array([u[1] for u in units])
    reported = np.array([u[2] for u in units])
    p = solve_fill_probability(rate, weights, eligible, reported)
    assert 0 <= p <= 1
    target = rate * weights[eligible].sum()
    reporting = weights[eligible & reported].sum()
    remaining = weights[eligible & ~reported].sum()
    expected = reporting + p * remaining
    if remaining <= 0:
        assert p == 0
    elif 0 < p < 1:
        assert np.isclose(expected, target, rtol=1e-9, atol=1e-6)
    elif p == 0:
        assert reporting >= target - 1e-6
    else:
        assert reporting + remaining <= target + 1e-6


@settings(max_examples=200, deadline=None)
@given(
    _units, st.floats(0, 1), st.floats(0, 1), st.floats(0, 1), st.integers(0, 2**32 - 1)
)
def test_reporters_claim_entitled_units_share_the_solved_probability_others_the_new_rate(
    units, rate, newly_entitled_rate, other_rate, seed
):
    weights = np.array([u[0] for u in units])
    entitled = np.array([u[1] for u in units])
    reported = np.array([u[2] for u in units])
    gb = np.ones(len(units), dtype=bool)
    draws = np.random.default_rng(seed).random(len(units))
    result, p = pension_credit_takeup_flags(
        draws, rate, weights, entitled, reported, gb, newly_entitled_rate
    )
    assert p == solve_fill_probability(rate, weights, entitled, reported)
    assert result[reported].all()
    entitled_non_reporters = entitled & ~reported
    np.testing.assert_array_equal(
        result[entitled_non_reporters], (draws < p)[entitled_non_reporters]
    )
    others = ~entitled & ~reported
    np.testing.assert_array_equal(result[others], (draws < newly_entitled_rate)[others])
    # The newly entitled rate changes neither the probability nor any flag of
    # an entitled unit, so the calibration year's take-up is unchanged.
    other, p_other = pension_credit_takeup_flags(
        draws, rate, weights, entitled, reported, gb, other_rate
    )
    assert p_other == p
    np.testing.assert_array_equal(other[entitled], result[entitled])


def test_weighted_take_up_among_eligible_matches_rate():
    rng = np.random.default_rng(0)
    n = 200_000
    weights = rng.uniform(100, 3_000, n)
    entitled = rng.random(n) < 0.2
    reported = entitled & (rng.random(n) < 0.4)
    result, _ = pension_credit_takeup_flags(
        rng.random(n), 0.62, weights, entitled, reported, np.ones(n, bool), 0.37
    )
    take_up = weights[entitled & result].sum() / weights[entitled].sum()
    assert abs(take_up - 0.62) < 0.005
    # Units with no entitlement claim at the newly entitled rate.
    assert abs(result[~entitled].mean() - 0.37) < 0.01


def test_rate_is_dwp_fye_2024_caseload_take_up():
    assert load_take_up_rate("pension_credit", 2025) == 0.62


def test_newly_entitled_rate_is_dwp_savings_credit_only_take_up():
    assert load_take_up_rate("pension_credit_newly_entitled", 2025) == 0.37
    assert load_take_up_rate("pension_credit_newly_entitled", 2022) == 0.42


def test_targets_reconcile_with_dwp_components():
    """Guarantee + savings credit spending and the three claim types sum to
    the transcribed totals (DWP Pension Credit sheet, separate rows)."""
    guarantee = {2022: 4_594.32, 2025: 5_787.44, 2030: 5_543.38}
    savings = {2022: 340.99, 2025: 357.03, 2030: 278.38}
    claims = {2022: (727, 448, 199), 2025: (807, 399, 176), 2030: (734, 285, 114)}
    targets = {t.name: t for t in dwp_pension_credit.get_targets()}
    spending = targets["dwp/pension_credit"]
    caseload = targets["dwp/pension_credit_claims"]
    for year in guarantee:
        assert abs(spending.values[year] / 1e6 - guarantee[year] - savings[year]) < 0.02
        assert abs(caseload.values[year] / 1e3 - sum(claims[year])) <= 1
    for target in targets.values():
        assert target.countries == GREAT_BRITAIN
        assert target.unit == (Unit.COUNT if target.is_count else Unit.GBP)


def test_pension_credit_is_targeted_only_by_dwp():
    names = {t.name for t in get_all_targets() if t.variable == "pension_credit"}
    assert names == {"dwp/pension_credit", "dwp/pension_credit_claims"}


def test_obr_pension_credit_row_is_not_parsed():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "4.9"
    ws["B2"] = "Pension credit"
    for col in "CDEFGHI":
        ws[f"{col}2"] = 6.0
    assert not obr._parse_welfare(wb)


def test_built_dataset_survey_reporters_claim_and_caseload_is_near_dwp(
    baseline, enhanced_frs
):
    """Every survey reporter claims. SPI-synthetic reports are imputed, so
    those units are drawn like non-reporters."""
    from policyengine_uk_data.datasets.pension_credit_takeup import (
        spi_synthetic_benunits,
    )
    from policyengine_uk_data.targets.build_loss_matrix import (
        _SimContext,
        restrict_to_countries,
    )

    year = 2025
    reported = (
        baseline.calculate("pension_credit_reported", year, map_to="benunit").values > 0
    ) & ~spi_synthetic_benunits(enhanced_frs)
    would_claim = baseline.calculate("would_claim_pc", year).values.astype(bool)
    assert would_claim[reported].all()

    baseline.default_calculation_period = str(year)
    ctx = _SimContext(baseline, str(year), None, None)
    weight = baseline.calculate("household_weight", year).values
    for target in dwp_pension_credit.get_targets():
        column = restrict_to_countries(
            target.custom_compute(ctx, target, year), ctx.country, target.countries
        )
        ratio = (column * weight).sum() / target.values[year]
        assert 0.5 < ratio < 2, (target.name, ratio)


def test_spi_synthetic_benefit_units_are_found_through_their_household():
    from types import SimpleNamespace

    import pandas as pd

    from policyengine_uk_data.datasets.pension_credit_takeup import (
        spi_synthetic_benunits,
    )

    dataset = SimpleNamespace(
        household=pd.DataFrame(
            {"household_id": [1, 2, 3], "household_is_spi_synthetic": [0, 1, 0]}
        ),
        person=pd.DataFrame(
            {
                "person_household_id": [1, 2, 2, 3],
                "person_benunit_id": [10, 20, 21, 30],
            }
        ),
        benunit=pd.DataFrame({"benunit_id": [30, 21, 20, 10]}),
    )
    np.testing.assert_array_equal(
        spi_synthetic_benunits(dataset), [False, True, True, False]
    )
    dataset.household = dataset.household.drop(columns="household_is_spi_synthetic")
    assert not spi_synthetic_benunits(dataset).any()


@settings(max_examples=200, deadline=None)
@given(_units, _units, st.floats(0, 1), st.integers(0, 2**32 - 1))
def test_northern_ireland_cannot_change_the_gb_solution(gb_units, ni_units, rate, seed):
    """DWP's take-up rate covers Great Britain, so Northern Ireland's
    entitlement, reporting and weights leave the probability and every GB
    flag unchanged."""

    def arrays(units):
        return (
            np.array([u[0] for u in units]),
            np.array([u[1] for u in units]),
            np.array([u[2] for u in units]),
        )

    gw, ge, gr = arrays(gb_units)
    nw, ne, nr = arrays(ni_units)
    draws = np.random.default_rng(seed).random(len(gw) + len(nw))
    gb_only, p_gb = pension_credit_takeup_flags(
        draws[: len(gw)], rate, gw, ge, gr, np.ones(len(gw), dtype=bool), 0.37
    )
    combined, p_uk = pension_credit_takeup_flags(
        draws,
        rate,
        np.concatenate([gw, nw]),
        np.concatenate([ge, ne]),
        np.concatenate([gr, nr]),
        np.concatenate([np.ones(len(gw), bool), np.zeros(len(nw), bool)]),
        0.37,
    )
    assert p_uk == p_gb
    np.testing.assert_array_equal(combined[: len(gw)], gb_only)
    # Entitled Northern Ireland non-reporters are drawn at the GB probability,
    # the rest at the newly entitled rate.
    ni_draws = draws[len(gw) :]
    ni_flags = combined[len(gw) :]
    np.testing.assert_array_equal(ni_flags[ne & ~nr], (ni_draws < p_gb)[ne & ~nr])
    np.testing.assert_array_equal(ni_flags[~ne & ~nr], (ni_draws < 0.37)[~ne & ~nr])
