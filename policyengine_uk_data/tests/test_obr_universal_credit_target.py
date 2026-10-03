"""Tests for the OBR universal credit target.

OBR EFO table 4.9 splits universal credit between spending inside the welfare
cap and outside it (the Intensive Work Search group, DWP's UC equivalent of
JSA). The split is not the household benefit cap, and policyengine-uk cannot
identify the Intensive Work Search group, so the two rows are one target:
GB universal credit, the sum of both rows.
"""

from functools import lru_cache
from types import SimpleNamespace

import numpy as np
import openpyxl
import pandas as pd
import pytest
from hypothesis import given, settings
from hypothesis import strategies as st

from policyengine_uk_data.storage import STORAGE_FOLDER
from policyengine_uk_data.targets import build_loss_matrix
from policyengine_uk_data.targets.build_loss_matrix import (
    _compute_column,
    restrict_to_countries,
)
from policyengine_uk_data.targets.schema import GREAT_BRITAIN, GeographicLevel
from policyengine_uk_data.targets.sources import obr

COUNTRIES = ("ENGLAND", "SCOTLAND", "WALES", "NORTHERN_IRELAND")
YEAR_COLUMNS = dict(zip("CDEFGHI", range(2024, 2031)))


def _committed_table() -> openpyxl.Workbook:
    return openpyxl.load_workbook(STORAGE_FOLDER / "obr_efo" / "efo_expenditure.xlsx")


def _uc_targets(wb) -> dict:
    return {t.name: t for t in obr._parse_welfare(wb) if "universal_credit" in t.name}


@lru_cache(maxsize=1)
def _committed_target():
    return _uc_targets(_committed_table())["obr/universal_credit"]


def test_universal_credit_is_one_gb_target():
    targets = _uc_targets(_committed_table())
    assert list(targets) == ["obr/universal_credit"]
    target = targets["obr/universal_credit"]
    assert target.variable == "universal_credit"
    assert target.countries == GREAT_BRITAIN
    # March 2026 EFO table 4.9, 2025-26: £66.411bn inside the welfare cap
    # (row 18) plus £12.876bn outside it (row 43).
    assert target.values[2025] == pytest.approx(79.287e9, abs=1e6)


def test_target_is_the_sum_of_both_rows_in_every_year():
    wb = _committed_table()
    ws = wb["4.9"]
    rows = [
        row
        for row in range(1, 56)
        if str(ws[f"B{row}"].value or "").startswith("Universal credit")
    ]
    assert rows == [18, 43]
    target = _uc_targets(wb)["obr/universal_credit"]
    assert sorted(target.values) == list(YEAR_COLUMNS.values())
    for column, year in YEAR_COLUMNS.items():
        expected = sum(ws[f"{column}{row}"].value for row in rows) * 1e9
        assert target.values[year] == pytest.approx(expected)


def _table(rows: list[tuple]) -> openpyxl.Workbook:
    """A minimal sheet 4.9: (label, value per year) in column B onwards."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "4.9"
    for i, (label, value) in enumerate(rows, start=6):
        ws[f"B{i}"] = label
        if value is not None:
            for column in YEAR_COLUMNS:
                ws[f"{column}{i}"] = value
    return wb


def test_rows_are_found_by_section_not_position():
    wb = _table(
        [
            ("Welfare cap", None),
            ("Pension credit", 6.0),
            ("Universal credit", 60.0),
            ("Child benefit", 13.0),
            ("Welfare spending outside the welfare cap", None),
            ("State pension", 140.0),
            ("Universal credit", 10.0),
        ]
    )
    target = _uc_targets(wb)["obr/universal_credit"]
    assert set(target.values.values()) == {70e9}


@pytest.mark.parametrize(
    "rows",
    [
        # No outside-the-cap row: a partial total would understate UC.
        [
            ("Universal credit", 60.0),
            ("Welfare spending outside the welfare cap", None),
            ("State pension", 140.0),
        ],
        # No section heading: the rows cannot be told apart.
        [("Universal credit", 60.0), ("Universal credit", 10.0)],
        # Two rows in one section.
        [
            ("Universal credit", 60.0),
            ("Universal credit", 1.0),
            ("Welfare spending outside the welfare cap", None),
            ("Universal credit", 10.0),
        ],
    ],
)
def test_no_target_unless_both_rows_are_unambiguous(rows):
    assert _uc_targets(_table(rows)) == {}


def test_no_target_when_the_rows_cover_different_years():
    wb = _table(
        [
            ("Universal credit", 60.0),
            ("Welfare spending outside the welfare cap", None),
            ("Universal credit", 10.0),
        ]
    )
    wb["4.9"]["I8"] = None  # 2030-31 missing from the outside-the-cap row only
    assert _uc_targets(wb) == {}


def test_rows_below_row_55_are_found():
    padding = [(f"Other benefit {i}", 1.0) for i in range(60)]
    wb = _table(
        [("Universal credit", 60.0)]
        + padding
        + [("Welfare spending outside the welfare cap", None)]
        + padding
        + [("Universal credit", 10.0)]
    )
    target = _uc_targets(wb)["obr/universal_credit"]
    assert set(target.values.values()) == {70e9}


def test_target_matrix_counts_gb_households_only(monkeypatch):
    """Through create_target_matrix itself: England has two benefit units on
    UC, Northern Ireland and Wales one each, Scotland none."""
    import policyengine_uk

    target = _committed_target()
    uc = pd.Series([100.0, 50.0, 300.0, 20.0])
    benunit_household = np.array([0, 0, 1, 2])
    country = pd.Series(["ENGLAND", "NORTHERN_IRELAND", "WALES", "SCOTLAND"])

    class FakeMicrosimulation:
        tax_benefit_system = SimpleNamespace(
            variables={
                "universal_credit": SimpleNamespace(
                    entity=SimpleNamespace(key="benunit")
                )
            }
        )

        def __init__(self, dataset=None, reform=None):
            pass

        def calculate(self, variable, *args, **kwargs):
            return {"universal_credit": uc, "country": country}[variable]

        def map_result(self, values, source, target_entity):
            assert (source, target_entity) == ("benunit", "household")
            return np.bincount(
                benunit_household, weights=np.asarray(values), minlength=len(country)
            )

    monkeypatch.setattr(policyengine_uk, "Microsimulation", FakeMicrosimulation)
    monkeypatch.setattr(
        build_loss_matrix,
        "get_all_targets",
        lambda geographic_level=None: (
            [target] if geographic_level == GeographicLevel.NATIONAL else []
        ),
    )

    matrix, values = build_loss_matrix.create_target_matrix(
        SimpleNamespace(time_period="2025"), time_period="2025"
    )
    np.testing.assert_array_equal(matrix["obr/universal_credit"], [150, 0, 20, 0])
    assert values["obr/universal_credit"] == target.values[2025]


def _fake_ctx(uc, benunit_household, household_country):
    n_households = len(household_country)
    variables = {
        "universal_credit": SimpleNamespace(entity=SimpleNamespace(key="benunit"))
    }
    return SimpleNamespace(
        sim=SimpleNamespace(
            tax_benefit_system=SimpleNamespace(variables=variables),
            calculate=lambda variable, *args, **kwargs: uc,
        ),
        household_from_family=lambda values: np.bincount(
            benunit_household, weights=np.asarray(values, float), minlength=n_households
        ),
        country=household_country,
    )


@settings(max_examples=200, deadline=None)
@given(
    st.lists(st.sampled_from(COUNTRIES), min_size=1, max_size=20).flatmap(
        lambda countries: st.tuples(
            st.just(np.array(countries)),
            st.lists(
                st.tuples(
                    st.integers(0, len(countries) - 1),
                    st.floats(0, 1e5, allow_nan=False),
                ),
                max_size=40,
            ),
        )
    )
)
def test_column_counts_gb_universal_credit_once(case):
    """Property: the column sums UC over benefit units in GB households, each
    once, and is zero for every Northern Ireland household."""
    household_country, benunits = case
    benunit_household = np.array([h for h, _ in benunits], dtype=int)
    uc = np.array([amount for _, amount in benunits], dtype=float)
    ctx = _fake_ctx(uc, benunit_household, household_country)
    target = _committed_target()

    column = restrict_to_countries(
        _compute_column(target, ctx, 2025), ctx.country, target.countries
    )

    in_gb = household_country[benunit_household] != "NORTHERN_IRELAND"
    assert column.sum() == pytest.approx(uc[in_gb].sum())
    assert (column[household_country == "NORTHERN_IRELAND"] == 0).all()
    assert (column >= 0).all()
