"""Tests for the DWP Housing Benefit targets by age group and GB scope."""

from types import SimpleNamespace

import numpy as np
import openpyxl
import pandas as pd
from hypothesis import given, settings
from hypothesis import strategies as st

from policyengine_uk_data.targets import get_all_targets
from policyengine_uk_data.targets.build_loss_matrix import restrict_to_countries
from policyengine_uk_data.targets.schema import GREAT_BRITAIN, Unit
from policyengine_uk_data.targets.sources import dwp_housing_benefit, obr

# Other rows of the same DWP Housing benefits sheet, transcribed separately:
# total Housing Benefit is AME within the welfare cap + AME outside it + LA
# funded (£ million), and the total caseload (thousands).
_DWP_TOTAL_GBP_M = {
    2022: 14_839.6 + 160.8 + 578.7,
    2023: 14_946.4 + 111.0 + 715.2,
    2024: 14_593.4 + 57.2 + 804.0,
    2025: 12_103.3 + 0.0 + 789.6,
    2026: 11_643.4 + 0.0 + 830.8,
    2027: 11_969.5 + 0.0 + 864.6,
    2028: 12_352.1 + 0.0 + 900.8,
    2029: 12_950.1 + 0.0 + 949.2,
    2030: 13_368.2 + 0.0 + 978.5,
}
_DWP_TOTAL_CASELOAD_K = {
    2022: 2_509,
    2023: 2_351,
    2024: 2_072,
    2025: 1_569,
    2026: 1_410,
    2027: 1_396,
    2028: 1_390,
    2029: 1_397,
    2030: 1_414,
}

_CALIBRATED = {
    "dwp/housing_benefit/over_pension_credit_age",
    "dwp/housing_benefit/over_pension_credit_age_claims",
}
_ALL = _CALIBRATED | {
    "dwp/housing_benefit/under_pension_credit_age",
    "dwp/housing_benefit/under_pension_credit_age_claims",
}


def _targets():
    """Both age groups, including the one calibration leaves out."""
    return {t.name: t for t in dwp_housing_benefit.build_targets()}


def test_targets_cover_great_britain_in_pounds_and_claims():
    targets = _targets()
    assert set(targets) == _ALL
    for name, target in targets.items():
        assert target.countries == GREAT_BRITAIN
        assert target.variable == "housing_benefit"
        assert target.is_count == name.endswith("_claims")
        assert target.unit == (Unit.COUNT if target.is_count else Unit.GBP)
    over = targets["dwp/housing_benefit/over_pension_credit_age"]
    under = targets["dwp/housing_benefit/under_pension_credit_age_claims"]
    assert over.values[2025] == 7_114.7e6
    assert under.values[2025] == 460e3


def test_age_split_reconciles_with_dwp_totals():
    """Over + under matches DWP's total, including LA-funded spending."""
    targets = _targets()
    over = targets["dwp/housing_benefit/over_pension_credit_age"].values
    under = targets["dwp/housing_benefit/under_pension_credit_age"].values
    over_k = targets["dwp/housing_benefit/over_pension_credit_age_claims"].values
    under_k = targets["dwp/housing_benefit/under_pension_credit_age_claims"].values
    for year, total in _DWP_TOTAL_GBP_M.items():
        assert abs((over[year] + under[year]) / 1e6 - total) < 0.25, year
    for year, total in _DWP_TOTAL_CASELOAD_K.items():
        # DWP rounds each caseload to the nearest thousand.
        assert abs((over_k[year] + under_k[year]) / 1e3 - total) <= 1, year


def test_only_pension_age_housing_benefit_is_calibrated():
    housing_benefit_targets = {
        t.name for t in get_all_targets() if t.variable == "housing_benefit"
    }
    assert housing_benefit_targets == _CALIBRATED


def test_obr_housing_benefit_row_is_not_parsed():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "4.9"
    ws["B2"] = "Housing benefit (not on JSA)1"
    ws["B3"] = "Pension credit"
    for row in (2, 3):
        for col in "CDEFGHI":
            ws[f"{col}{row}"] = 6.0
    names = {t.name for t in obr._parse_welfare(wb)}
    assert "obr/pension_credit" in names
    assert not any("housing_benefit" in name for name in names)


def _ctx(
    benunit_of_person,
    is_adult,
    is_sp_age,
    benunit_hb,
    household_of_benunit,
    benunit_benefits=None,
):
    """Minimal loss-matrix context over explicit entity mappings.

    ``benunit_benefits`` maps benefit-unit variables such as
    ``universal_credit`` to amounts; any other benefit is zero.
    """
    benunit_of_person = np.asarray(benunit_of_person)
    household_of_benunit = np.asarray(household_of_benunit)
    n_benunits = len(benunit_hb)
    n_households = household_of_benunit.max() + 1 if n_benunits else 0

    def map_result(values, source, target):
        values = np.asarray(values, dtype=float)
        if (source, target) == ("person", "benunit"):
            return np.bincount(benunit_of_person, values, minlength=n_benunits)
        if (source, target) == ("benunit", "household"):
            return np.bincount(household_of_benunit, values, minlength=n_households)
        raise AssertionError((source, target))

    person = {"is_adult": np.asarray(is_adult), "is_SP_age": np.asarray(is_sp_age)}
    benunit_values = {"housing_benefit": benunit_hb, **(benunit_benefits or {})}
    sim = SimpleNamespace(
        map_result=map_result,
        calculate=lambda variable: SimpleNamespace(
            values=np.asarray(
                benunit_values.get(variable, np.zeros(n_benunits)), dtype=float
            )
        ),
    )
    return SimpleNamespace(
        sim=sim,
        pe_person=lambda variable: person[variable],
        household_from_family=lambda values: map_result(values, "benunit", "household"),
    )


def _column(ctx, name):
    target = _targets()[name]
    return target.custom_compute(ctx, target, 2025)


def test_benefit_rules_assign_mixed_age_couples():
    # Benefit units, one household each: pensioner couple; mixed-age couple
    # on pension-age rules; working-age single; pensioner with an
    # 18-year-old dependant; mixed-age couple whose younger partner gets
    # income-related ESA; mixed-age couple on Universal Credit.
    ctx = _ctx(
        benunit_of_person=[0, 0, 1, 1, 2, 3, 3, 4, 4, 5, 5],
        is_adult=[1] * 11,
        is_sp_age=[1, 1, 1, 0, 0, 1, 0, 1, 0, 1, 0],
        benunit_hb=[5_000, 4_000, 3_000, 2_000, 1_000, 500],
        household_of_benunit=[0, 1, 2, 3, 4, 5],
        benunit_benefits={
            "esa_income": [0, 0, 0, 0, 4_000, 0],
            "universal_credit": [0, 0, 0, 0, 0, 6_000],
        },
    )
    over = _column(ctx, "dwp/housing_benefit/over_pension_credit_age")
    under = _column(ctx, "dwp/housing_benefit/under_pension_credit_age")
    np.testing.assert_array_equal(over, [5_000, 4_000, 0, 2_000, 0, 0])
    np.testing.assert_array_equal(under, [0, 0, 3_000, 0, 1_000, 500])


@st.composite
def _population(draw):
    n_benunits = draw(st.integers(1, 12))
    n_households = draw(st.integers(1, n_benunits))
    household_of_benunit = draw(
        st.lists(
            st.integers(0, n_households - 1),
            min_size=n_benunits,
            max_size=n_benunits,
        )
    )
    n_people = draw(st.integers(n_benunits, 4 * n_benunits))
    # Every benefit unit has at least one member.
    benunit_of_person = list(range(n_benunits)) + draw(
        st.lists(
            st.integers(0, n_benunits - 1),
            min_size=n_people - n_benunits,
            max_size=n_people - n_benunits,
        )
    )
    flags = st.lists(st.booleans(), min_size=n_people, max_size=n_people)
    benunit_hb = draw(
        st.lists(
            st.one_of(st.just(0.0), st.floats(1, 50_000)),
            min_size=n_benunits,
            max_size=n_benunits,
        )
    )
    benefit = st.lists(
        st.sampled_from([0.0, 100.0]), min_size=n_benunits, max_size=n_benunits
    )
    return dict(
        benunit_of_person=benunit_of_person,
        is_adult=draw(flags),
        is_sp_age=draw(flags),
        benunit_hb=benunit_hb,
        household_of_benunit=household_of_benunit,
        benunit_benefits={
            name: draw(benefit) for name in dwp_housing_benefit._WORKING_AGE_BENEFITS
        },
    )


@settings(max_examples=200, deadline=None)
@given(_population())
def test_age_split_partitions_housing_benefit(population):
    """Over + under is all Housing Benefit, in pounds and in claims, and
    the older group is exactly the units with an adult over SPA and no
    working-age income benefit."""
    ctx = _ctx(**population)
    household = np.asarray(population["household_of_benunit"])
    hb = np.asarray(population["benunit_hb"])
    n_households = household.max() + 1
    over = _column(ctx, "dwp/housing_benefit/over_pension_credit_age")
    under = _column(ctx, "dwp/housing_benefit/under_pension_credit_age")
    over_k = _column(ctx, "dwp/housing_benefit/over_pension_credit_age_claims")
    under_k = _column(ctx, "dwp/housing_benefit/under_pension_credit_age_claims")
    np.testing.assert_allclose(
        over + under, np.bincount(household, hb, minlength=n_households)
    )
    np.testing.assert_array_equal(
        over_k + under_k, np.bincount(household, hb > 0, minlength=n_households)
    )
    people = pd.DataFrame(
        {
            "benunit": population["benunit_of_person"],
            "older": np.asarray(population["is_adult"])
            & np.asarray(population["is_sp_age"]),
        }
    )
    older_unit = (
        people.groupby("benunit").older.any().reindex(range(len(hb)), fill_value=False)
    ).to_numpy()
    on_working_age_benefit = np.any(
        [np.asarray(v) > 0 for v in population["benunit_benefits"].values()], axis=0
    )
    expected_over = np.bincount(
        household, hb * (older_unit & ~on_working_age_benefit), minlength=n_households
    )
    np.testing.assert_allclose(over, expected_over)


_COUNTRIES = ["ENGLAND", "SCOTLAND", "WALES", "NORTHERN_IRELAND"]


@settings(max_examples=200, deadline=None)
@given(
    st.lists(
        st.tuples(st.floats(-1e9, 1e9), st.sampled_from(_COUNTRIES)),
        min_size=1,
        max_size=50,
    )
)
def test_country_restriction_partitions_the_uk(rows):
    column = np.array([value for value, _ in rows])
    country = np.array([c for _, c in rows])
    gb = restrict_to_countries(column, country, GREAT_BRITAIN)
    ni = restrict_to_countries(column, country, ("NORTHERN_IRELAND",))
    np.testing.assert_allclose(gb + ni, column)
    np.testing.assert_array_equal(gb[country == "NORTHERN_IRELAND"], 0)
    np.testing.assert_array_equal(restrict_to_countries(gb, country, GREAT_BRITAIN), gb)
    assert restrict_to_countries(column, country, None) is column
