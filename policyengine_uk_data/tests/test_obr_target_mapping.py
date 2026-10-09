"""Pin which PE-UK variable each OBR receipts target calibrates.

The cash-receipts parser once targeted total NICs ("National insurance
contributions", Table 3.8) on ``ni_employee``, the variable the NICs parser
targets at Class 1 employee NICs alone. These tests read the committed EFO
workbooks offline and pin the receipts targets' variables, the rows the NIC
targets come from, and why total NICs has no target of its own.
"""

from unittest.mock import patch

import openpyxl
import pytest
import requests

from policyengine_uk_data.storage import STORAGE_FOLDER
from policyengine_uk_data.targets.sources import obr

# Every target parsed from the receipts workbook. Expenditure-side and static
# OBR targets are left to their own tests.
EXPECTED_VARIABLES = {
    "obr/income_tax": "income_tax",
    "obr/vat": "vat",
    "obr/fuel_duties": "fuel_duty",
    "obr/capital_gains_tax": "capital_gains_tax",
    "obr/sdlt": "stamp_duty_land_tax",
    "obr/ni_employee": "ni_employee",
    "obr/ni_employer": "ni_employer",
    "obr/ni_self_employed": "ni_self_employed",
}

# Table 3.4 rows under "National insurance contributions" in the committed
# workbook: the class targets, then the rows the model leaves empty.
CLASS_ROWS = {
    "obr/ni_employee": "Class 1 Employee NICs",
    "obr/ni_employer": "Class 1 Employer NICs",
    "obr/ni_self_employed": "Class 4 and Class 2 Self employed NICs",
}
UNMODELLED_ROWS = ("Statutory payment recoveries", "Other NIC")


@pytest.fixture(scope="module")
def offline_targets():
    obr._download_workbook.cache_clear()

    def get(*args, **kwargs):
        raise requests.ConnectionError("offline")

    with (
        patch.object(obr.requests, "get", side_effect=get),
        patch.object(obr.time, "sleep", lambda s: None),
    ):
        targets = obr.get_targets()
    obr._download_workbook.cache_clear()
    return {t.name: t for t in targets}


@pytest.fixture(scope="module")
def receipts():
    return openpyxl.load_workbook(
        STORAGE_FOLDER / "obr_efo" / "efo_receipts.xlsx", data_only=True
    )


def _row_2025(ws, label: str, column: str) -> float:
    """Value for 2025-26 (in £) of the row whose label starts with ``label``."""
    return ws[f"{column}{obr._find_row(ws, label)}"].value * 1e9


def test_receipts_target_variables_are_pinned():
    wb = obr._fallback_workbook("efo-receipts")
    targets = obr._parse_receipts(wb) + obr._parse_nics(wb)
    assert {t.name: t.variable for t in targets} == EXPECTED_VARIABLES


def test_total_nics_row_exists_but_has_no_target(offline_targets, receipts):
    """The absence comes from the mapping, not from a missing row."""
    cash = obr._find_receipts_sheet(receipts)
    total = _row_2025(cash, "National insurance contributions", "E")
    assert total > 150e9
    assert "obr/ni" not in offline_targets
    assert all(t.values.get(2025) != total for t in offline_targets.values())


@pytest.mark.parametrize("name,label", CLASS_ROWS.items())
def test_nic_class_targets_read_table_3_4(offline_targets, receipts, name, label):
    expected = _row_2025(receipts["3.4"], label, "D")
    assert offline_targets[name].values[2025] == pytest.approx(expected)


def test_class_targets_and_unmodelled_rows_make_up_total_nics(receipts):
    """Why total NICs is not targeted on total_national_insurance.

    The class targets plus rows the model leaves empty (statutory payment
    recoveries; Class 1A, 1B and 3 in "Other NIC") sum to the Table 3.4 total,
    so a total target would demand NICs no household carries.
    """
    ws = receipts["3.4"]
    total = _row_2025(ws, "National insurance contributions", "D")
    classes = sum(_row_2025(ws, label, "D") for label in CLASS_ROWS.values())
    unmodelled = sum(_row_2025(ws, label, "D") for label in UNMODELLED_ROWS)
    assert unmodelled > 0
    assert classes + unmodelled == pytest.approx(total, abs=0.02e9)
