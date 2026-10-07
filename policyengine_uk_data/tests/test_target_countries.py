"""Tests for the countries a calibration target covers.

DWP's statistics cover Great Britain, and England and Wales for the benefits
devolved to the Scottish Government. A target that sets ``countries`` counts
only households in those countries in its loss matrix column.
"""

from itertools import combinations
from types import SimpleNamespace

import numpy as np
import openpyxl
import pandas as pd
import pytest
from hypothesis import given, settings
from hypothesis import strategies as st
from policyengine_uk.variables.household.demographic.country import Country

from policyengine_uk_data.storage import STORAGE_FOLDER
from policyengine_uk_data.targets import build_loss_matrix
from policyengine_uk_data.targets.build_loss_matrix import restrict_to_countries
from policyengine_uk_data.targets.registry import discover_source_modules
from policyengine_uk_data.targets.schema import (
    ENGLAND_AND_WALES,
    GREAT_BRITAIN,
    GeographicLevel,
    Target,
    Unit,
)
from policyengine_uk_data.targets.sources import (
    dwp,
    dwp_housing_benefit,
    dwp_pension_credit,
    hmrc_salary_sacrifice,
    obr,
    ons_labour_market,
)

COUNTRIES = tuple(country.name for country in Country)
YEAR_COLUMNS = dict(zip("CDEFGHI", range(2024, 2031)))
# Benefits whose executive competence passed to the Scottish Government, so
# DWP reports them for England and Wales (BECT Notes K, L and N).
DEVOLVED_IN_SCOTLAND = {
    "pip",
    "attendance_allowance",
    "carers_allowance",
    "winter_fuel_allowance",
}


def test_coverage_constants_are_values_of_the_country_variable():
    # A misspelt country would silently zero a column.
    assert set(GREAT_BRITAIN) == {"ENGLAND", "SCOTLAND", "WALES"}
    assert set(ENGLAND_AND_WALES) == {"ENGLAND", "WALES"}
    assert set(GREAT_BRITAIN) <= set(COUNTRIES)


# ── restrict_to_countries ──────────────────────────────────────────────


@st.composite
def _households(draw):
    """A household column and each household's country, of equal length."""
    n = draw(st.integers(0, 30))
    country = draw(st.lists(st.sampled_from(COUNTRIES), min_size=n, max_size=n))
    column = draw(
        st.lists(st.floats(-1e12, 1e12, allow_nan=False), min_size=n, max_size=n)
    )
    return np.array(column), np.array(country, dtype=object)


_country_sets = st.sets(st.sampled_from(COUNTRIES)).map(tuple)


@settings(max_examples=300, deadline=None)
@given(_households(), _country_sets)
def test_restriction_keeps_households_in_scope_and_zeroes_the_rest(case, countries):
    column, country = case
    restricted = restrict_to_countries(column, country, countries)
    in_scope = np.array([c in countries for c in country], dtype=bool)
    np.testing.assert_array_equal(restricted[in_scope], column[in_scope])
    assert (restricted[~in_scope] == 0).all()


@settings(max_examples=300, deadline=None)
@given(_households(), _country_sets)
def test_restrictions_to_complementary_countries_add_up_to_the_column(case, countries):
    """Conservation: splitting the UK into two sets of countries loses and
    double counts nothing, so a GB column plus its Northern Ireland (and
    unknown-country) remainder is the UK column."""
    column, country = case
    rest = tuple(c for c in COUNTRIES if c not in countries)
    np.testing.assert_array_equal(
        restrict_to_countries(column, country, countries)
        + restrict_to_countries(column, country, rest),
        column,
    )


@settings(max_examples=300, deadline=None)
@given(_households(), _country_sets, _country_sets)
def test_restricting_twice_restricts_to_the_common_countries(case, first, second):
    column, country = case
    common = tuple(c for c in first if c in second)
    np.testing.assert_array_equal(
        restrict_to_countries(
            restrict_to_countries(column, country, first), country, second
        ),
        restrict_to_countries(column, country, common),
    )


@settings(max_examples=100, deadline=None)
@given(_households())
def test_no_countries_means_the_whole_uk(case):
    column, country = case
    assert restrict_to_countries(column, country, None) is column


def test_target_matrix_counts_only_households_in_a_targets_countries(monkeypatch):
    """Through create_target_matrix, on a real simulation with one household in
    each country, so the coverage constants meet the `country` variable's own
    output. Each country's household holds a distinct power of two, so a
    column's total says which countries it counted."""
    import policyengine_uk

    amount = {"ENGLAND": 1.0, "SCOTLAND": 2.0, "WALES": 4.0, "NORTHERN_IRELAND": 8.0}
    situation = {
        "people": {c: {"age": {2025: 40}} for c in amount},
        "benunits": {c: {"members": [c]} for c in amount},
        "households": {c: {"members": [c], "country": {2025: c}} for c in amount},
    }
    microsimulation = policyengine_uk.Microsimulation

    def target(name, countries):
        return Target(
            name=name,
            variable="universal_credit",
            source="test",
            unit=Unit.GBP,
            values={2025: 1.0},
            countries=countries,
            custom_compute=lambda ctx, target, year: np.array(
                [amount[c] for c in ctx.country]
            ),
        )

    targets = [
        target("uk", None),
        target("gb", GREAT_BRITAIN),
        target("ew", ENGLAND_AND_WALES),
        target("ni", ("NORTHERN_IRELAND",)),
    ]
    monkeypatch.setattr(
        policyengine_uk,
        "Microsimulation",
        lambda dataset=None, reform=None: microsimulation(situation=situation),
    )
    monkeypatch.setattr(
        build_loss_matrix,
        "get_all_targets",
        lambda geographic_level=None: (
            targets if geographic_level == GeographicLevel.NATIONAL else []
        ),
    )

    matrix, values = build_loss_matrix.create_target_matrix(
        SimpleNamespace(time_period="2025"), time_period="2025"
    )
    assert list(matrix.columns) == ["uk", "gb", "ew", "ni"]
    assert matrix.sum().to_dict() == {"uk": 15, "gb": 7, "ew": 5, "ni": 8}
    assert list(values) == [1.0] * 4


# ── OBR table 4.9 ──────────────────────────────────────────────────────


@pytest.fixture(scope="module")
def table_4_9():
    wb = openpyxl.load_workbook(STORAGE_FOLDER / "obr_efo" / "efo_expenditure.xlsx")
    return wb, wb["4.9"]


def _label(ws, row) -> str:
    return str(ws[f"B{row}"].value or "").strip()


def _dwp_blocks(ws) -> list[range]:
    """Table 4.9's "DWP social security" blocks, inside and outside the
    welfare cap: from each block's total to the "Other DWP" line closing it."""
    blocks, start = [], None
    for row in range(1, ws.max_row + 1):
        if _label(ws, row).startswith("DWP social security"):
            start = row
        elif start is not None and _label(ws, row).startswith("Other DWP"):
            blocks.append(range(start, row + 1))
            start = None
    return blocks


def _source_rows(ws, target) -> list[tuple[int, ...]]:
    """The table rows a target's values come from: one row, or two summed."""

    def matches(rows):
        return all(
            sum(ws[f"{column}{row}"].value for row in rows) * 1e9
            == pytest.approx(target.values[year], rel=1e-12)
            for column, year in YEAR_COLUMNS.items()
            if year in target.values
        )

    numeric = [
        row
        for row in range(1, ws.max_row + 1)
        if all(
            isinstance(ws[f"{column}{row}"].value, (int, float))
            for column in YEAR_COLUMNS
        )
    ]
    singles = [(row,) for row in numeric if matches((row,))]
    return singles or [pair for pair in combinations(numeric, 2) if matches(pair)]


def test_dwp_blocks_hold_exactly_the_dwp_social_security_lines(table_4_9):
    """Each block's lines add up to its "DWP social security" total in every
    year, so the blocks hold DWP's lines and nothing else. Northern Ireland's
    benefits are separate "NI social security" rows outside the blocks."""
    _, ws = table_4_9
    blocks = _dwp_blocks(ws)
    assert len(blocks) == 2
    for block in blocks:
        total_row, lines = block[0], block[1:]
        for column in YEAR_COLUMNS:
            values = [ws[f"{column}{row}"].value for row in lines]
            numbers = [v for v in values if isinstance(v, (int, float))]
            # "*" marks a line under £0.1bn.
            tolerance = 0.1 * values.count("*") + 0.005
            assert sum(numbers) == pytest.approx(
                ws[f"{column}{total_row}"].value, abs=tolerance
            )
    ni_rows = [
        row
        for row in range(1, ws.max_row + 1)
        if _label(ws, row).startswith("NI social security")
    ]
    assert len(ni_rows) == 2
    assert not any(row in block for row in ni_rows for block in blocks)


def test_welfare_targets_cover_the_countries_of_their_table_block(table_4_9):
    """Lines in a "DWP social security" block cover Great Britain, or England
    and Wales for benefits devolved to Scotland; lines outside one (child
    benefit, HMRC's) cover the UK."""
    wb, ws = table_4_9
    blocks = _dwp_blocks(ws)
    targets = obr._parse_welfare(wb)
    assert {"obr/state_pension", "obr/child_benefit", "obr/pip"} <= {
        t.name for t in targets
    }
    for target in targets:
        sources = _source_rows(ws, target)
        assert len(sources) == 1, (target.name, sources)
        in_dwp_block = [any(row in block for block in blocks) for row in sources[0]]
        if target.countries is None:
            assert not any(in_dwp_block), target.name
            continue
        assert all(in_dwp_block), target.name
        expected = (
            ENGLAND_AND_WALES
            if target.variable in DEVOLVED_IN_SCOTLAND
            else GREAT_BRITAIN
        )
        assert target.countries == expected, target.name


# ── DWP statistics ─────────────────────────────────────────────────────


def test_every_dwp_target_declares_the_countries_it_covers():
    """No DWP statistic covers Northern Ireland, so a DWP target without
    ``countries`` would count Northern Ireland households."""
    modules = [
        module
        for module in discover_source_modules()
        if module.__name__.rsplit(".", 1)[-1].startswith("dwp")
    ]
    targets = [target for module in modules for target in module.get_targets()]
    assert len(targets) > 100
    undeclared = [t.name for t in targets if t.countries is None]
    assert undeclared == []
    assert all(set(t.countries) <= set(GREAT_BRITAIN) for t in targets)


def test_pip_claimant_targets_cover_england_and_wales():
    """Adult Disability Payment replaced PIP in Scotland; DWP's PIP
    statistics cover England and Wales."""
    pip = [t for t in dwp.get_targets() if t.name.startswith("dwp/pip_")]
    assert len(pip) == 2
    assert all(t.countries == ENGLAND_AND_WALES for t in pip)


# ── Targets introduced or replaced in the 1.58.0 batch ──────────────────


@pytest.mark.parametrize(
    "get_targets, names",
    [
        (
            dwp_housing_benefit.get_targets,
            {
                "dwp/housing_benefit/over_pension_credit_age",
                "dwp/housing_benefit/over_pension_credit_age_claims",
                "dwp/housing_benefit/under_pension_credit_age_general_needs",
                "dwp/housing_benefit/under_pension_credit_age_general_needs_claims",
            },
        ),
        (
            dwp_housing_benefit.build_targets,
            {
                "dwp/housing_benefit/over_pension_credit_age",
                "dwp/housing_benefit/over_pension_credit_age_claims",
                "dwp/housing_benefit/under_pension_credit_age",
                "dwp/housing_benefit/under_pension_credit_age_claims",
                "dwp/housing_benefit/under_pension_credit_age_general_needs",
                "dwp/housing_benefit/under_pension_credit_age_general_needs_claims",
            },
        ),
        (
            dwp_pension_credit.get_targets,
            {"dwp/pension_credit", "dwp/pension_credit_claims"},
        ),
    ],
    ids=[
        "housing-benefit-calibration",
        "housing-benefit-diagnostics",
        "pension-credit",
    ],
)
def test_batch_dwp_replacement_targets_cover_exactly_great_britain(get_targets, names):
    """#490 Housing Benefit and #510 Pension Credit use DWP's GB tables.

    Equality catches losing Scotland or Wales as well as adding Northern
    Ireland; a GB-subset check would miss the former.
    """
    targets = get_targets()
    assert len(targets) == len(names)
    assert {t.name: t.countries for t in targets} == dict.fromkeys(names, GREAT_BRITAIN)


def test_batch_obr_replacements_preserve_one_gb_universal_credit_target(monkeypatch):
    """#530 combines both DWP GB UC rows; #490, #510 and #533 replace
    obsolete OBR targets with DWP or HMRC targets rather than duplicating them.
    """
    monkeypatch.setattr(obr, "_download_workbook", obr._fallback_workbook)
    targets = obr.get_targets()
    names = [t.name for t in targets]
    assert len(names) == len(set(names))
    uc = [t for t in targets if t.variable == "universal_credit"]
    assert [(t.name, t.countries) for t in uc] == [
        ("obr/universal_credit", GREAT_BRITAIN)
    ]
    assert not {
        "obr/housing_benefit",
        "obr/pension_credit",
        "obr/universal_credit_in_cap",
        "obr/universal_credit_outside_cap",
        "obr/salary_sacrifice_employee_ni_relief",
        "obr/salary_sacrifice_employer_ni_relief",
    } & set(names)


def test_lfs_employment_targets_keep_the_whole_uk_in_scope():
    """#529's MGRN and MGRQ source series explicitly cover the UK."""
    targets = ons_labour_market.get_targets()
    assert len(targets) == 2
    assert {t.name: t.countries for t in targets} == {
        "ons/lfs_employees": None,
        "ons/lfs_self_employed": None,
    }


def test_hmrc_salary_sacrifice_replacements_preserve_mains_country_scope(monkeypatch):
    """#533's HMRC targets keep main's unrestricted country scope.

    The committed table does not give a country field, so this refresh
    leaves main's scope intact. It must not inherit a DWP GB restriction.
    """
    monkeypatch.setattr(
        hmrc_salary_sacrifice,
        "_read_table",
        lambda url: pd.read_csv(
            hmrc_salary_sacrifice.FALLBACK_CSV, dtype=str, encoding="utf-8-sig"
        ),
    )
    targets = hmrc_salary_sacrifice.get_targets()
    names = {
        "hmrc/salary_sacrifice_it_relief_basic_rate",
        "hmrc/salary_sacrifice_it_relief_higher_rate",
        "hmrc/salary_sacrifice_it_relief_additional_rate",
        "hmrc/salary_sacrifice_employee_nics_relief",
        "hmrc/salary_sacrifice_employer_nics_relief",
        "hmrc/salary_sacrifice_contributions",
    }
    assert len(targets) == len(names)
    assert {t.name: t.countries for t in targets} == dict.fromkeys(names)


def test_uc_payment_distribution_targets_all_cover_great_britain():
    """The batch's repaired open top bands share the extract's GB scope."""
    targets = dwp._uc_payment_distribution_targets()
    assert targets
    assert all(t.countries == GREAT_BRITAIN for t in targets)
    assert {t.name for t in targets if not np.isfinite(t.upper_bound)} == {
        f"dwp/uc_payment_dist/{family_type}_annual_payment_30_000_to_inf"
        for family_type in (
            "SINGLE",
            "LONE_PARENT",
            "COUPLE_NO_CHILDREN",
            "COUPLE_WITH_CHILDREN",
        )
    }
