"""Hash-bound OBR Economic and fiscal outlook workbooks.

Each EFO vintage the OBR targets are built from has a manifest in
``targets/source_manifests/``. It names the two "detailed forecast tables"
workbooks (receipts and expenditure) with their SHA-256, size, title and
provenance. It also gives the unit and the financial years every parsed
table must carry, and the targets the vintage must yield. ``sources.yaml``
names the manifest in use.

Builds read only the committed bytes, verified against the manifest, and
never download. A target's ``forecast_vintage`` therefore always names the
publication its bytes came from. Year columns are read from each table's
header row, so a vintage whose tables start a year later is parsed by
year rather than by column letter. A missing row, year or target fails the
build instead of being logged and dropped.

Adopting a new vintage (for example on Budget day):

    python -m policyengine_uk_data.targets.obr_sources ingest \\
        --vintage october_2026 --label "October 2026" \\
        --receipts receipts.xlsx --expenditure expenditure.xlsx \\
        --first-year 2025 --last-year 2031 --dry-run

obr.uk refuses scripted downloads (HTTP 403 to GitHub Actions runners, as
11e41f9d records, and to a developer machine on 2026-10-15), so a person
downloads the two "detailed forecast tables" workbooks in a browser and
passes their paths. ``--dry-run`` checks everything and prints each
workbook's title, SHA-256, size and download time and every target's old
and new value, writing nothing. It also says whether the bytes match the
committed workbooks when the vintage is the active one, which makes a
March 2026 download a rehearsal. Without ``--dry-run``, ``ingest`` parses
the workbooks in a staging folder and then copies them under
``storage/obr_efo/<vintage>/``. It requires each workbook's title to name
the vintage, and requires the targets the current manifest requires. It
writes the manifest with each file's SHA-256, size and download time (the
file's modification time unless ``--downloaded-at`` is given) and never
replaces an ingested vintage. Point ``sources.yaml`` at the new manifest to
adopt it.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import re
import sys
from dataclasses import dataclass
from datetime import date, datetime, timezone
from functools import lru_cache
from pathlib import Path
from typing import Mapping, Optional, Sequence

import openpyxl
import yaml

MANIFEST_DIR = Path(__file__).parent / "source_manifests"
KINDS = ("receipts", "expenditure")
UNIT_SCALES = {"£ billion": 1e9}
_FINANCIAL_YEAR = re.compile(r"^\s*(\d{4})\s*[-–]\s*(\d{2})\s*$")
_HEADER_ROWS = 10
_UNIT_ROWS = 4


class SourceError(Exception):
    """A source file, table or target set does not match its manifest.

    Not a ValueError, so a parser's "row not found" handling cannot swallow
    it: a table without the declared years or unit stops the build.
    """


@dataclass(frozen=True)
class SourceFile:
    kind: str
    path: str  # relative to the storage folder
    sha256: str
    size_bytes: int
    title: str  # the workbook's own title, cell B2 of its Contents sheet
    url: str
    retrieved: str
    # When the file was downloaded (ISO 8601 with a UTC offset), if known.
    downloaded_at: Optional[str] = None


@dataclass(frozen=True)
class ObrManifest:
    vintage: str
    label: str
    first_year: int  # financial years by the calendar year they start in
    last_year: int
    unit: str
    files: Mapping[str, SourceFile]
    required: tuple  # target names
    alternatives: tuple  # each a tuple of sets of names; one set must be present
    path: Optional[Path] = None

    @property
    def years(self) -> list:
        return list(range(self.first_year, self.last_year + 1))

    @property
    def scale(self) -> float:
        return UNIT_SCALES[self.unit]


def parse_manifest(raw: Mapping, path: Optional[Path] = None) -> ObrManifest:
    where = str(path or "manifest")
    if (raw or {}).get("schema_version") != 1:
        raise SourceError(f"{where}: schema_version must be 1")
    unit = raw.get("unit")
    if unit not in UNIT_SCALES:
        raise SourceError(f"{where}: unit must be one of {sorted(UNIT_SCALES)}")
    first, last = int(raw["first_year"]), int(raw["last_year"])
    if not first <= last:
        raise SourceError(f"{where}: first_year is after last_year")
    files = {}
    for kind in KINDS:
        entry = (raw.get("files") or {}).get(kind)
        if not entry:
            raise SourceError(f"{where}: no {kind} workbook")
        sha = str(entry["sha256"])
        if not re.fullmatch(r"[0-9a-f]{64}", sha):
            raise SourceError(f"{where}: {kind} sha256 must be 64 hex characters")
        files[kind] = SourceFile(
            kind=kind,
            path=str(entry["path"]),
            sha256=sha,
            size_bytes=int(entry["size_bytes"]),
            title=str(entry["title"]),
            url=str(entry["url"]),
            retrieved=str(entry["retrieved"]),
            downloaded_at=_timestamp(entry.get("downloaded_at"), f"{where}: {kind}"),
        )
    required, alternatives = [], []
    for item in (raw.get("targets") or {}).get("required") or []:
        if isinstance(item, str):
            required.append(item)
        elif isinstance(item, Mapping) and item.get("any_of"):
            alternatives.append(tuple(frozenset(group) for group in item["any_of"]))
        else:
            raise SourceError(f"{where}: unreadable required target {item!r}")
    if len(set(required)) != len(required):
        raise SourceError(f"{where}: a required target is listed twice")
    if not required:
        raise SourceError(f"{where}: no required targets")
    vintage, label = str(raw["vintage"]), str(raw["label"])
    check_names(vintage, label, {kind: f.title for kind, f in files.items()}, where)
    return ObrManifest(
        vintage=str(raw["vintage"]),
        label=str(raw["label"]),
        first_year=first,
        last_year=last,
        unit=unit,
        files=files,
        required=tuple(required),
        alternatives=tuple(alternatives),
        path=path,
    )


def _timestamp(value, where: str) -> Optional[str]:
    """An ISO 8601 time with a UTC offset, or None."""
    if value is None:
        return None
    text = value.isoformat() if isinstance(value, datetime) else str(value)
    try:
        parsed = datetime.fromisoformat(text)
    except ValueError as error:
        raise SourceError(
            f"{where}: downloaded_at {text!r} is not an ISO 8601 time"
        ) from error
    if parsed.tzinfo is None:
        raise SourceError(f"{where}: downloaded_at {text!r} needs a UTC offset")
    return parsed.isoformat()


def check_names(vintage: str, label: str, titles: Mapping[str, str], where: str):
    """The vintage, its label and each workbook's own title must agree.

    ``march_2026`` must be labelled "March 2026", and each workbook titled
    "March 2026 Economic and fiscal outlook ..." ending with its kind. The
    title is read from the verified bytes, so the bytes decide the vintage.
    """
    if not re.fullmatch(r"[a-z]+_\d{4}", vintage):
        raise SourceError(f"{where}: vintage must look like october_2026")
    if label.lower().replace(" ", "_") != vintage:
        raise SourceError(f"{where}: label {label!r} does not name vintage {vintage!r}")
    for kind, title in titles.items():
        lowered = title.strip().lower()
        if not lowered.startswith(f"{label.lower()} economic and fiscal outlook") or (
            not lowered.endswith(kind)
        ):
            raise SourceError(
                f"{where}: the {kind} workbook is titled {title!r}, which is not "
                f"the {label} EFO {kind} tables"
            )


def load_manifest(path) -> ObrManifest:
    path = Path(path)
    if not path.is_absolute():
        path = MANIFEST_DIR / path
    with open(path) as f:
        return parse_manifest(yaml.safe_load(f), path)


@lru_cache(maxsize=1)
def active_manifest() -> ObrManifest:
    """The manifest ``sources.yaml`` names for the OBR targets."""
    from policyengine_uk_data.targets.sources._common import load_config

    return load_manifest(load_config()["obr"]["manifest"])


def _storage() -> Path:
    from policyengine_uk_data.storage import STORAGE_FOLDER

    return Path(STORAGE_FOLDER)


def sha256_bytes(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def verified_bytes(entry: SourceFile, storage: Optional[Path] = None) -> bytes:
    """The file's bytes, if and only if they are the ones the manifest names."""
    path = (storage or _storage()) / entry.path
    if not path.is_file():
        raise SourceError(f"OBR {entry.kind} workbook {path} is missing")
    data = path.read_bytes()
    if len(data) != entry.size_bytes or sha256_bytes(data) != entry.sha256:
        raise SourceError(
            f"OBR {entry.kind} workbook {path} is not the file its manifest "
            f"names: {len(data)} bytes, sha256 {sha256_bytes(data)}, against "
            f"{entry.size_bytes} bytes, sha256 {entry.sha256}"
        )
    return data


def workbook_title(wb: openpyxl.Workbook) -> str:
    """The title an EFO workbook gives itself (Contents sheet, cell B2)."""
    if "Contents" not in wb.sheetnames:
        raise SourceError("the workbook has no Contents sheet")
    return str(wb["Contents"]["B2"].value or "").strip()


def load_workbook(
    kind: str,
    manifest: Optional[ObrManifest] = None,
    storage: Optional[Path] = None,
) -> openpyxl.Workbook:
    """Open a manifest's verified workbook, checking it carries its own title."""
    manifest = manifest or active_manifest()
    entry = manifest.files[kind]
    wb = openpyxl.load_workbook(
        io.BytesIO(verified_bytes(entry, storage)), data_only=False
    )
    if workbook_title(wb) != entry.title:
        raise SourceError(
            f"OBR {kind} workbook is titled {workbook_title(wb)!r}, not {entry.title!r}"
        )
    return wb


def year_columns(ws, first_year: int, last_year: int) -> dict:
    """Map the column letters of a table's financial-year headers to years.

    The header row is the first of the top rows with cells reading like
    "2025-26"; each maps to the calendar year it starts in. The years must
    be exactly ``first_year`` to ``last_year``, each once, so a table that
    starts a year later, or misses a year, fails rather than being read
    with the wrong labels.
    """
    for row in ws.iter_rows(min_row=1, max_row=_HEADER_ROWS):
        found = {}
        for cell in row:
            match = (
                _FINANCIAL_YEAR.match(cell.value)
                if isinstance(cell.value, str)
                else None
            )
            if not match:
                continue
            start, end = int(match[1]), int(match[2])
            if (start + 1) % 100 != end:
                raise SourceError(
                    f"{ws.title}!{cell.coordinate}: {cell.value!r} is not one "
                    "financial year"
                )
            if start in found.values():
                raise SourceError(
                    f"{ws.title}: {cell.value!r} heads two columns of row {cell.row}"
                )
            found[cell.column_letter] = start
        if found:
            years = sorted(found.values())
            if years != list(range(first_year, last_year + 1)):
                raise SourceError(
                    f"{ws.title}: header row {row[0].row} covers "
                    f"{', '.join(str(y) for y in years)}, but the manifest "
                    f"declares {first_year} to {last_year}"
                )
            return found
    raise SourceError(
        f"{ws.title}: no row of financial-year headers in rows 1-{_HEADER_ROWS}"
    )


def unit_scale(ws, unit: str) -> float:
    """The multiplier for a table's values, from its unit label."""
    for row in ws.iter_rows(min_row=1, max_row=_UNIT_ROWS):
        for cell in row:
            if isinstance(cell.value, str) and cell.value.strip().startswith(unit):
                return UNIT_SCALES[unit]
    raise SourceError(f"{ws.title}: no {unit!r} unit label in rows 1-{_UNIT_ROWS}")


# A row label giving its own unit, as table 4.19's "Standard colour licence
# fee (£)" or "Number of UK households (m)" do under a "£ billion unless
# otherwise stated" table.
_ROW_UNIT = re.compile(
    r"\(\s*(?:£|£m|£bn|£ ?million|£ ?billion|m|bn|million|thousands?|per cent|%|p|pence)\s*\)"
    r"|\bper cent\b|%"
)


def read_row(ws, row: int, manifest: ObrManifest) -> dict:
    """A table row's numeric values by year, in pounds.

    Refuses a row whose label states a unit of its own: the table's unit is
    applied to every row read.
    """
    label = str(ws[f"B{row}"].value or "")
    if _ROW_UNIT.search(label):
        raise SourceError(
            f"{ws.title}!B{row}: {label!r} states its own unit, not the table's "
            f"{manifest.unit}"
        )
    scale = unit_scale(ws, manifest.unit)
    out = {}
    for column, year in year_columns(
        ws, manifest.first_year, manifest.last_year
    ).items():
        value = ws[f"{column}{row}"].value
        if isinstance(value, (int, float)) and not isinstance(value, bool):
            out[year] = float(value) * scale
    return out


def coverage_problems(targets: Sequence, manifest: ObrManifest) -> list:
    """Everything wrong with the targets built from a vintage's workbooks.

    ``targets`` are all the targets the workbooks yielded. Each must carry
    the manifest's vintage; each required target must be present with every
    year of the manifest; one set of each ``any_of`` group must be present;
    and no workbook target may be missing from the manifest, so the manifest
    stays the complete list.
    """
    built = {t.name: t for t in targets}
    problems = [
        f"{t.name} carries vintage {t.forecast_vintage!r}, not {manifest.vintage!r}"
        for t in targets
        if t.forecast_vintage != manifest.vintage
    ]
    for name in manifest.required:
        if name not in built:
            problems.append(f"{name} was not built")
    for groups in manifest.alternatives:
        if not any(group <= set(built) for group in groups):
            options = " or ".join("+".join(sorted(g)) for g in groups)
            problems.append(f"none of {options} was built")
    declared = set(manifest.required).union(
        *(group for groups in manifest.alternatives for group in groups)
    )
    for name, target in sorted(built.items()):
        if name not in declared:
            problems.append(f"{name} was built but is not in the manifest")
        missing = [y for y in manifest.years if y not in target.values]
        if missing:
            problems.append(
                f"{name} has no value for {', '.join(str(y) for y in missing)}"
            )
    return problems


def check_coverage(targets: Sequence, manifest: ObrManifest) -> None:
    problems = coverage_problems(targets, manifest)
    if problems:
        raise SourceError(
            f"OBR {manifest.label} targets are incomplete: " + "; ".join(problems)
        )


def target_diff(old: Sequence, new: Sequence) -> list:
    """Each workbook target's value by year in two target sets."""
    before = {t.name: t for t in old}
    after = {t.name: t for t in new}
    rows = []
    for name in sorted(set(before) | set(after)):
        a, b = before.get(name), after.get(name)
        years = sorted(set(a.values if a else {}) | set(b.values if b else {}))
        for year in years:
            x = a.values.get(year) if a else None
            y = b.values.get(year) if b else None
            rows.append({"target": name, "year": year, "old": x, "new": y})
    return rows


def _manifest_yaml(manifest: ObrManifest, alternatives_raw) -> str:
    raw = {
        "schema_version": 1,
        "vintage": manifest.vintage,
        "label": manifest.label,
        "first_year": manifest.first_year,
        "last_year": manifest.last_year,
        "unit": manifest.unit,
        "files": {
            kind: {
                "path": f.path,
                "sha256": f.sha256,
                "size_bytes": f.size_bytes,
                "title": f.title,
                "url": f.url,
                "retrieved": f.retrieved,
                **({"downloaded_at": f.downloaded_at} if f.downloaded_at else {}),
            }
            for kind, f in manifest.files.items()
        },
        "targets": {
            "required": list(manifest.required)
            + [{"any_of": groups} for groups in alternatives_raw]
        },
    }
    header = (
        f"# OBR Economic and fiscal outlook, {manifest.label}: detailed forecast\n"
        "# tables. Written by `python -m policyengine_uk_data.targets.obr_sources\n"
        "# ingest`; see that module for what each field means.\n"
    )
    return header + yaml.safe_dump(raw, sort_keys=False, allow_unicode=True)


def ingest(
    *,
    vintage: str,
    label: str,
    sources: Mapping[str, Path],
    urls: Mapping[str, str],
    first_year: int,
    last_year: int,
    current: Optional[ObrManifest] = None,
    storage: Optional[Path] = None,
    manifest_dir: Optional[Path] = None,
    retrieved: Optional[str] = None,
    downloaded_at: Optional[Mapping[str, str]] = None,
    dry_run: bool = False,
):
    """Adopt a vintage's two workbooks. Returns (manifest, manifest path, diff).

    Each workbook's SHA-256, size and download time go into the manifest.
    The download time is ``downloaded_at[kind]`` if given, otherwise the
    file's modification time, which a browser sets when it saves a download.
    With ``dry_run`` everything is checked and the diff returned, but
    nothing is written; the manifest returned is not on disk.

    Everything is checked in a staging folder first: both workbooks read
    and carry the vintage's titles, and the build's parsers yield every
    required target and year from them. Only then are the files copied
    into storage and the manifest written. A vintage already in storage or
    in the manifests folder is never replaced.
    """
    import tempfile

    from policyengine_uk_data.targets.sources import obr

    current = current or active_manifest()
    storage = Path(storage or _storage())
    manifest_dir = Path(manifest_dir or MANIFEST_DIR)
    retrieved = retrieved or (
        f"downloaded in a browser and ingested on {date.today().isoformat()}"
    )
    downloaded_at = dict(downloaded_at or {})
    destination = storage / "obr_efo" / vintage
    path = manifest_dir / f"obr_{vintage}.yaml"
    if not dry_run and (destination.exists() or path.exists()):
        raise SourceError(
            f"{vintage} is already ingested ({destination} or {path} exists); "
            "remove it deliberately before ingesting again"
        )
    files, titles, data = {}, {}, {}
    for kind in KINDS:
        data[kind] = Path(sources[kind]).read_bytes()
        try:
            wb = openpyxl.load_workbook(io.BytesIO(data[kind]), data_only=False)
        except Exception as error:
            raise SourceError(
                f"{sources[kind]} is not an xlsx workbook ({type(error).__name__}: "
                f"{error}); obr.uk serves an HTML page to scripted downloads, so "
                "download the workbook in a browser"
            ) from error
        titles[kind] = workbook_title(wb)
        files[kind] = SourceFile(
            kind=kind,
            path=f"obr_efo/{vintage}/{kind}.xlsx",
            sha256=sha256_bytes(data[kind]),
            size_bytes=len(data[kind]),
            title=titles[kind],
            url=urls.get(kind, ""),
            retrieved=retrieved,
            downloaded_at=_timestamp(
                downloaded_at.get(kind)
                or datetime.fromtimestamp(
                    Path(sources[kind]).stat().st_mtime, tz=timezone.utc
                ),
                f"ingest: {kind}",
            ),
        )
    check_names(vintage, label, titles, "ingest")
    manifest = ObrManifest(
        vintage=vintage,
        label=label,
        first_year=first_year,
        last_year=last_year,
        unit=current.unit,
        files=files,
        required=current.required,
        alternatives=current.alternatives,
    )
    with tempfile.TemporaryDirectory() as staging:
        staging = Path(staging)
        for kind in KINDS:
            staged = staging / files[kind].path
            staged.parent.mkdir(parents=True, exist_ok=True)
            staged.write_bytes(data[kind])
        # Raises unless the new workbooks yield every target and year required.
        new = obr.workbook_targets(manifest, staging)
    old = obr.workbook_targets(current)
    text = _manifest_yaml(
        manifest,
        [[sorted(group) for group in groups] for groups in current.alternatives],
    )
    if dry_run:
        return manifest, path, target_diff(old, new)
    destination.mkdir(parents=True)
    for kind in KINDS:
        (storage / files[kind].path).write_bytes(data[kind])
    path.write_text(text)
    return load_manifest(path), path, target_diff(old, new)


def _gbp(value) -> str:
    return "" if value is None else f"{value / 1e9:,.3f}"


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(
        prog="python -m policyengine_uk_data.targets.obr_sources"
    )
    sub = parser.add_subparsers(dest="command", required=True)
    i = sub.add_parser("ingest", help="adopt a new EFO vintage's workbooks")
    i.add_argument("--vintage", required=True, help="e.g. october_2026")
    i.add_argument("--label", required=True, help='e.g. "October 2026"')
    i.add_argument("--receipts", required=True, type=Path)
    i.add_argument("--expenditure", required=True, type=Path)
    i.add_argument("--receipts-url", default="")
    i.add_argument("--expenditure-url", default="")
    i.add_argument("--first-year", required=True, type=int)
    i.add_argument("--last-year", required=True, type=int)
    i.add_argument("--retrieved", help="how the files were obtained")
    i.add_argument(
        "--downloaded-at",
        help="when the workbooks were downloaded (ISO 8601 with offset); "
        "default: each file's modification time",
    )
    i.add_argument(
        "--dry-run",
        action="store_true",
        help="check everything and print the change, but write nothing",
    )
    v = sub.add_parser("verify", help="verify the active manifest's workbooks")
    v.add_argument("--manifest", help="manifest file (default: the active one)")
    args = parser.parse_args(argv)
    if args.command == "verify":
        from policyengine_uk_data.targets.sources import obr

        manifest = load_manifest(args.manifest) if args.manifest else active_manifest()
        targets = obr.workbook_targets(manifest)
        check_coverage(targets, manifest)
        print(f"{manifest.label}: both workbooks verified, {len(targets)} targets")
        return 0
    manifest, path, diff = ingest(
        vintage=args.vintage,
        label=args.label,
        sources={"receipts": args.receipts, "expenditure": args.expenditure},
        urls={"receipts": args.receipts_url, "expenditure": args.expenditure_url},
        first_year=args.first_year,
        last_year=args.last_year,
        retrieved=args.retrieved,
        downloaded_at={kind: args.downloaded_at for kind in KINDS}
        if args.downloaded_at
        else None,
        dry_run=args.dry_run,
    )
    for kind, f in manifest.files.items():
        print(
            f"{kind}: {f.title}; sha256 {f.sha256}, {f.size_bytes:,} bytes, "
            f"downloaded {f.downloaded_at}"
        )
    print(f"{'would write' if args.dry_run else 'wrote'} {path}")
    print("target,year,current £bn,new £bn")
    for row in diff:
        print(f"{row['target']},{row['year']},{_gbp(row['old'])},{_gbp(row['new'])}")
    if not args.dry_run:
        print(f"To adopt it, set obr.manifest in sources.yaml to {path.name}.")
    else:
        current = active_manifest()
        if manifest.vintage == current.vintage:
            for kind, f in manifest.files.items():
                same = f.sha256 == current.files[kind].sha256
                print(
                    f"{kind}: {'the same bytes as' if same else 'different bytes from'} "
                    f"the committed {current.label} workbook"
                )
    return 0


if __name__ == "__main__":
    sys.exit(main())
